from __future__ import annotations
import csv
import io
import secrets
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.db import get_db
from app.models import CategoryNode, Equipment, InventoryItem, Site, Ticket, TicketWorkSession, User
from app.security import current_user
from app.access import has_permission
from app.routes.web import ctx, templates, forbidden
from app.services.audit import audit
from app.services.categories import category_options, category_path, build_category_tree, next_category_sort_order
from app.services.time_tracking import apply_session_cost, recalculate_ticket_labor_cost, format_duration, money

router = APIRouter()


def _user(request: Request, db: Session):
    return current_user(request, db)


def _norm_header(value) -> str:
    s = str(value or '').strip().lower().replace('ё', 'е')
    for ch in ' -./\\()№':
        s = s.replace(ch, '_')
    while '__' in s:
        s = s.replace('__', '_')
    return s.strip('_')


def _rows_from_upload(upload: UploadFile) -> list[dict[str, object]]:
    raw = upload.file.read(8 * 1024 * 1024 + 1)
    if len(raw) > 8 * 1024 * 1024:
        raise ValueError('Файл больше 8 МБ')
    suffix = Path(upload.filename or '').suffix.lower()
    rows: list[dict[str, object]] = []
    if suffix == '.xlsx':
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            return []
        keys = [_norm_header(x) for x in header]
        for values in iterator:
            if not any(v not in (None, '') for v in values):
                continue
            rows.append({keys[i]: values[i] if i < len(values) else None for i in range(len(keys)) if keys[i]})
        return rows
    if suffix in {'.csv', '.txt'}:
        text = None
        for enc in ('utf-8-sig', 'utf-8', 'cp1251'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise ValueError('Не удалось определить кодировку CSV')
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=';,\t,')
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ';'
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        for row in reader:
            normalized = {_norm_header(k): v for k, v in row.items() if k}
            if any(str(v or '').strip() for v in normalized.values()):
                rows.append(normalized)
        return rows
    raise ValueError('Поддерживаются файлы .xlsx и .csv')


def _pick(row: dict, *keys: str, default=''):
    for key in keys:
        value = row.get(_norm_header(key))
        if value not in (None, ''):
            return value
    return default


def _as_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _as_float(value, default=0.0):
    if value in (None, ''):
        return default
    try:
        return float(str(value).strip().replace(' ', '').replace(',', '.'))
    except Exception:
        return default


def _as_money(value) -> Decimal:
    try:
        return Decimal(str(value or 0).strip().replace(' ', '').replace(',', '.')).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError, AttributeError):
        return Decimal('0.00')


def _template_xlsx(headers: list[str], examples: list[list[object]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in examples:
        ws.append(row)
    fill = PatternFill('solid', fgColor='172235')
    for cell in ws[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, 1):
        width = max(14, min(32, len(header) + 5))
        ws.column_dimensions[get_column_letter(idx)].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.get('/data-import', response_class=HTMLResponse)
def data_import_page(request: Request, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'data.import'):
        return forbidden(request, db, u, 'data.import')
    return templates.TemplateResponse('data_import.html', ctx(request, db, result=None))


@router.get('/data-import/equipment-template.xlsx')
def equipment_import_template(request: Request, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u or not has_permission(u, 'data.import'):
        return Response(status_code=403)
    headers = ['inventory_no','name','site','address','category','model','serial_no','status','criticality','parent_inventory_no','owner_username','installed_at','warranty_until','next_maintenance_at']
    example = [['ХО-000173','Бонета морозильная №7','Магазин №1','Петропавловск','Холодильное оборудование / Бонеты','Demo Frost 2500','FRZ-24-00918','working','high','','tech','2025-03-01','2027-03-01','2026-10-15']]
    data = _template_xlsx(headers, example, 'CMDB')
    return Response(data, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename="FMTS_CMDB_import_template.xlsx"'})


@router.get('/data-import/inventory-template.xlsx')
def inventory_import_template(request: Request, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u or not has_permission(u, 'data.import'):
        return Response(status_code=403)
    headers = ['sku','name','qty','min_qty','unit','unit_cost']
    example = [['ZIP-0001','Компрессор холодильный',3,1,'шт',45000]]
    data = _template_xlsx(headers, example, 'ЗИП')
    return Response(data, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename="FMTS_inventory_import_template.xlsx"'})


@router.post('/data-import/equipment', response_class=HTMLResponse)
def equipment_import(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'data.import'):
        return forbidden(request, db, u, 'data.import')
    result = {'kind':'equipment','created':0,'updated':0,'skipped':0,'errors':[]}
    try:
        rows = _rows_from_upload(file)
        pending_parents = []
        status_map = {'работает':'working','рабочее':'working','неисправно':'broken','неисправен':'broken','в ремонте':'repair','ремонт':'repair','списано':'decommissioned','списан':'decommissioned'}
        valid_statuses = {'working','broken','repair','decommissioned'}
        valid_criticality = {'low','normal','high','critical'}
        for n, row in enumerate(rows, 2):
            inv = str(_pick(row,'inventory_no','инвентарный_номер','инв_номер','инвентарный №')).strip()
            name = str(_pick(row,'name','наименование','оборудование')).strip()
            site_name = str(_pick(row,'site','объект','магазин','подразделение')).strip()
            if not inv or not name or not site_name:
                result['skipped'] += 1
                result['errors'].append(f'Строка {n}: нужны inventory_no, name и site')
                continue
            site = db.query(Site).filter(Site.name == site_name).first()
            address = str(_pick(row,'address','адрес')).strip()
            if not site:
                site = Site(name=site_name, address=address)
                db.add(site); db.flush()
            elif address and not site.address:
                site.address = address
            obj = db.query(Equipment).filter(Equipment.inventory_no == inv).first()
            is_new = obj is None
            if is_new:
                obj = Equipment(site_id=site.id, name=name, inventory_no=inv, qr_token=secrets.token_urlsafe(24))
                db.add(obj)
            obj.site_id = site.id
            obj.name = name
            obj.category = str(_pick(row,'category','категория', default=obj.category or 'Прочее')).strip() or 'Прочее'
            obj.model = str(_pick(row,'model','модель', default=obj.model or '')).strip()
            obj.serial_no = str(_pick(row,'serial_no','серийный_номер','серийный', default=obj.serial_no or '')).strip()
            raw_status = str(_pick(row,'status','статус', default=obj.status or 'working')).strip().lower()
            obj.status = status_map.get(raw_status, raw_status if raw_status in valid_statuses else 'working')
            crit = str(_pick(row,'criticality','критичность', default=obj.criticality or 'normal')).strip().lower()
            obj.criticality = crit if crit in valid_criticality else 'normal'
            obj.installed_at = _as_date(_pick(row,'installed_at','дата_установки')) or obj.installed_at
            obj.warranty_until = _as_date(_pick(row,'warranty_until','гарантия_до')) or obj.warranty_until
            obj.next_maintenance_at = _as_date(_pick(row,'next_maintenance_at','следующее_то')) or obj.next_maintenance_at
            owner = str(_pick(row,'owner_username','ответственный','владелец')).strip()
            if owner:
                owner_obj = db.query(User).filter((User.username == owner) | (User.email == owner)).first()
                if owner_obj:
                    obj.owner_user_id = owner_obj.id
            db.flush()
            pending_parents.append((obj.id, str(_pick(row,'parent_inventory_no','родитель','родительский_инвентарный_номер')).strip(), n))
            result['created' if is_new else 'updated'] += 1
        db.flush()
        for obj_id, parent_inv, n in pending_parents:
            if not parent_inv:
                continue
            obj = db.get(Equipment, obj_id)
            parent = db.query(Equipment).filter(Equipment.inventory_no == parent_inv).first()
            if parent and parent.id != obj.id:
                obj.parent_id = parent.id
            else:
                result['errors'].append(f'Строка {n}: родитель {parent_inv} не найден')
        db.commit()
        audit(db, request, u, 'cmdb.import', entity_type='equipment', details=f"created={result['created']}; updated={result['updated']}; skipped={result['skipped']}")
    except Exception as exc:
        db.rollback()
        result['errors'].append(str(exc))
    return templates.TemplateResponse('data_import.html', ctx(request, db, result=result))


@router.post('/data-import/inventory', response_class=HTMLResponse)
def inventory_import(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'data.import'):
        return forbidden(request, db, u, 'data.import')
    result = {'kind':'inventory','created':0,'updated':0,'skipped':0,'errors':[]}
    try:
        rows = _rows_from_upload(file)
        for n, row in enumerate(rows, 2):
            sku = str(_pick(row,'sku','артикул','код')).strip()
            name = str(_pick(row,'name','наименование')).strip()
            if not sku or not name:
                result['skipped'] += 1
                result['errors'].append(f'Строка {n}: нужны sku и name')
                continue
            obj = db.query(InventoryItem).filter(InventoryItem.sku == sku).first()
            is_new = obj is None
            if is_new:
                obj = InventoryItem(sku=sku, name=name)
                db.add(obj)
            obj.name = name
            obj.qty = _as_float(_pick(row,'qty','остаток','количество'), obj.qty or 0)
            obj.min_qty = _as_float(_pick(row,'min_qty','минимальный_остаток','мин_остаток'), obj.min_qty or 0)
            obj.unit = str(_pick(row,'unit','единица','ед_изм', default=obj.unit or 'шт')).strip() or 'шт'
            obj.unit_cost = _as_money(_pick(row,'unit_cost','цена','стоимость', default=obj.unit_cost or 0))
            result['created' if is_new else 'updated'] += 1
        db.commit()
        audit(db, request, u, 'inventory.import', entity_type='inventory', details=f"created={result['created']}; updated={result['updated']}; skipped={result['skipped']}")
    except Exception as exc:
        db.rollback()
        result['errors'].append(str(exc))
    return templates.TemplateResponse('data_import.html', ctx(request, db, result=result))


def _would_cycle(db: Session, node_id: int, parent_id: int | None) -> bool:
    seen = set()
    cur = db.get(CategoryNode, parent_id) if parent_id else None
    while cur:
        if cur.id == node_id or cur.id in seen:
            return True
        seen.add(cur.id)
        cur = db.get(CategoryNode, cur.parent_id) if cur.parent_id else None
    return False


@router.get('/categories', response_class=HTMLResponse)
def categories_page(request: Request, kind: str = 'ticket', db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'settings.manage'):
        return forbidden(request, db, u, 'settings.manage')
    if kind not in {'ticket','equipment','knowledge','service'}:
        kind = 'ticket'
    rows = db.query(CategoryNode).filter(CategoryNode.kind == kind).order_by(CategoryNode.sort_order, CategoryNode.name, CategoryNode.id).all()
    paths = {x.id: category_path(db, x) for x in rows}
    tree = build_category_tree(rows)
    return templates.TemplateResponse('categories.html', ctx(request, db, rows=rows, tree=tree, paths=paths, kind=kind, parents=category_options(db, kind, True)))


@router.post('/categories/new')
def category_new(request: Request, kind: str = Form('ticket'), name: str = Form(...), parent_id: str = Form(''), sort_order: str = Form(''), db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'settings.manage'):
        return forbidden(request, db, u, 'settings.manage')
    kind = kind if kind in {'ticket','equipment','knowledge','service'} else 'ticket'
    name = name.strip()
    parent = db.get(CategoryNode, int(parent_id)) if parent_id else None
    if parent and parent.kind != kind:
        parent = None
    parent_value = parent.id if parent else None
    if name and not db.query(CategoryNode).filter(CategoryNode.kind == kind, CategoryNode.name == name, CategoryNode.parent_id == parent_value).first():
        try:
            order_value = int(sort_order) if str(sort_order).strip() else next_category_sort_order(db, kind, parent_value)
        except Exception:
            order_value = next_category_sort_order(db, kind, parent_value)
        row = CategoryNode(kind=kind, name=name, code=name.lower().replace(' ','_')[:180], parent_id=parent_value, sort_order=order_value, active=True)
        db.add(row); db.commit(); db.refresh(row)
        audit(db, request, u, 'category.create', entity_type='category', entity_id=row.id, details=category_path(db,row))
    return RedirectResponse(f'/categories?kind={kind}', 303)


@router.post('/categories/{category_id}/update')
def category_update(category_id: int, request: Request, name: str = Form(...), parent_id: str = Form(''), sort_order: str = Form(''), active: str = Form('1'), db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'settings.manage'):
        return forbidden(request, db, u, 'settings.manage')
    row = db.get(CategoryNode, category_id)
    if not row:
        return RedirectResponse('/categories', 303)
    requested_parent = int(parent_id) if parent_id else None
    parent = db.get(CategoryNode, requested_parent) if requested_parent else None
    if parent and (parent.kind != row.kind or _would_cycle(db, row.id, parent.id)):
        parent = None
    old_parent_id = row.parent_id
    row.name = name.strip() or row.name
    row.parent_id = parent.id if parent else None
    if str(sort_order).strip():
        try:
            row.sort_order = int(sort_order)
        except Exception:
            pass
    elif old_parent_id != row.parent_id:
        row.sort_order = next_category_sort_order(db, row.kind, row.parent_id)
    row.active = active == '1'
    db.commit()
    audit(db, request, u, 'category.update', entity_type='category', entity_id=row.id, details=category_path(db,row))
    return RedirectResponse(f'/categories?kind={row.kind}', 303)


@router.post('/categories/{category_id}/toggle')
def category_toggle(category_id: int, request: Request, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'settings.manage'):
        return forbidden(request, db, u, 'settings.manage')
    row = db.get(CategoryNode, category_id)
    if not row:
        return RedirectResponse('/categories', 303)
    row.active = not bool(row.active)
    db.commit()
    audit(db, request, u, 'category.toggle', entity_type='category', entity_id=row.id, details=f"active={row.active}; path={category_path(db,row)}")
    return RedirectResponse(f'/categories?kind={row.kind}', 303)


@router.get('/reports/labor', response_class=HTMLResponse)
def labor_report(request: Request, days: int = 30, user_id: str = '', site_id: str = '', db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'labor.report'):
        return forbidden(request, db, u, 'labor.report')
    days = max(1, min(int(days or 30), 3660))
    since = datetime.utcnow() - timedelta(days=days)
    q = db.query(TicketWorkSession).filter(TicketWorkSession.started_at >= since, TicketWorkSession.ended_at.is_not(None))
    if user_id:
        q = q.filter(TicketWorkSession.user_id == int(user_id))
    if site_id:
        q = q.join(Ticket, Ticket.id == TicketWorkSession.ticket_id).filter(Ticket.site_id == int(site_id))
    entries = q.order_by(TicketWorkSession.started_at.desc()).limit(3000).all()
    grouped = {}
    total_seconds = 0
    total_amount = Decimal('0.00')
    for e in entries:
        seconds = int(e.duration_seconds or 0)
        amount = money(e.labor_amount)
        total_seconds += seconds
        total_amount += amount
        row = grouped.setdefault(e.user_id, {'user':e.user,'seconds':0,'amount':Decimal('0.00'),'sessions':0,'tickets':set()})
        row['seconds'] += seconds; row['amount'] += amount; row['sessions'] += 1; row['tickets'].add(e.ticket_id)
    summary = []
    for row in grouped.values():
        row['ticket_count'] = len(row.pop('tickets'))
        summary.append(row)
    summary.sort(key=lambda x: (-x['seconds'], x['user'].full_name if x['user'] else ''))
    return templates.TemplateResponse('labor_report.html', ctx(request, db, entries=entries, summary=summary, total_seconds=total_seconds, total_amount=total_amount, days=days, user_id=user_id, site_id=site_id, technicians=db.query(User).filter(User.role=='technician').order_by(User.full_name).all(), sites=db.query(Site).order_by(Site.name).all()))


@router.get('/reports/labor.csv')
def labor_report_csv(request: Request, days: int = 30, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u or not has_permission(u, 'labor.report'):
        return Response(status_code=403)
    since = datetime.utcnow() - timedelta(days=max(1,min(int(days or 30),3660)))
    entries = db.query(TicketWorkSession).filter(TicketWorkSession.started_at >= since, TicketWorkSession.ended_at.is_not(None)).order_by(TicketWorkSession.started_at).all()
    out = io.StringIO(); w = csv.writer(out, delimiter=';')
    w.writerow(['Дата','Сотрудник','Заявка','Объект','Минуты','Ставка, ₸/ч','Стоимость, ₸','Комментарий'])
    for e in entries:
        t = e.ticket
        w.writerow([e.started_at.strftime('%d.%m.%Y %H:%M'), e.user.full_name if e.user else '', t.number if t else '', t.site.name if t and t.site else '', round((e.duration_seconds or 0)/60,1), str(e.hourly_rate_snapshot or 0), str(e.labor_amount or 0), e.note or ''])
    payload = ('\ufeff' + out.getvalue()).encode('utf-8')
    return Response(payload, media_type='text/csv; charset=utf-8', headers={'Content-Disposition':'attachment; filename="FMTS_labor.csv"'})


@router.post('/reports/labor/recalculate')
def labor_recalculate(request: Request, db: Session = Depends(get_db)):
    u = _user(request, db)
    if not u:
        return RedirectResponse('/login', 303)
    if not has_permission(u, 'labor.rate.manage'):
        return forbidden(request, db, u, 'labor.rate.manage')
    rows = db.query(TicketWorkSession).filter(TicketWorkSession.ended_at.is_not(None), TicketWorkSession.labor_amount.is_(None)).all()
    ticket_ids = set()
    for e in rows:
        e.hourly_rate_snapshot = money(e.user.hourly_rate if e.user else 0)
        apply_session_cost(e)
        ticket_ids.add(e.ticket_id)
    for ticket_id in ticket_ids:
        recalculate_ticket_labor_cost(db, ticket_id)
    db.commit()
    audit(db, request, u, 'labor.recalculate', entity_type='work_session', details=f'sessions={len(rows)}; tickets={len(ticket_ids)}')
    return RedirectResponse('/reports/labor', 303)
